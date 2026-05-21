import argparse
import io
import tempfile
from datetime import datetime
from pathlib import Path


def parse_args():
    parser = argparse.ArgumentParser(
        description="Generate a PowerPoint deck from source-code files."
    )
    parser.add_argument(
        "sources",
        nargs="+",
        help="One or more source files or folders. Folders are scanned recursively for common code/text files.",
    )
    parser.add_argument(
        "--output",
        required=True,
        help="Output .pptx path.",
    )
    parser.add_argument(
        "--title",
        default="Code Review Deck",
        help="Presentation title.",
    )
    parser.add_argument(
        "--subtitle",
        default="Generated from local project files",
        help="Presentation subtitle.",
    )
    parser.add_argument(
        "--max-lines",
        type=int,
        default=34,
        help="Maximum number of code lines per slide.",
    )
    parser.add_argument(
        "--font-size",
        type=int,
        default=18,
        help="Code font size used in rendered images.",
    )
    parser.add_argument(
        "--style",
        default="monokai",
        help="Pygments style name for syntax highlighting.",
    )
    parser.add_argument(
        "--project-root",
        default=".",
        help="Project root used to shorten displayed file paths.",
    )
    return parser.parse_args()


def ensure_dependencies():
    missing = []
    try:
        from pptx import Presentation  # noqa: F401
        from pptx.dml.color import RGBColor  # noqa: F401
        from pptx.enum.text import PP_ALIGN  # noqa: F401
        from pptx.util import Inches, Pt  # noqa: F401
    except Exception:
        missing.append("python-pptx")
    try:
        from pygments import highlight  # noqa: F401
        from pygments.formatters import ImageFormatter  # noqa: F401
        from pygments.lexers import TextLexer, get_lexer_for_filename  # noqa: F401
    except Exception:
        missing.append("pygments")
    try:
        from PIL import Image  # noqa: F401
    except Exception:
        missing.append("pillow")
    if missing:
        raise RuntimeError(
            "Missing required packages: {}. Install them with: pip install {}".format(
                ", ".join(missing),
                " ".join(missing),
            )
        )


def expand_sources(source_items):
    allowed_suffixes = {
        ".py",
        ".ps1",
        ".js",
        ".ts",
        ".tsx",
        ".jsx",
        ".json",
        ".yaml",
        ".yml",
        ".md",
        ".txt",
        ".sql",
        ".html",
        ".css",
        ".xml",
    }
    resolved = []
    for item in source_items:
        path = Path(item).expanduser().resolve()
        if not path.exists():
            raise RuntimeError("Source path not found: {}".format(path))
        if path.is_file():
            resolved.append(path)
            continue
        for child in sorted(path.rglob("*")):
            if child.is_file() and child.suffix.lower() in allowed_suffixes:
                resolved.append(child.resolve())
    if not resolved:
        raise RuntimeError("No source files were found in the provided inputs.")
    return resolved


def chunk_lines(lines, max_lines):
    if max_lines <= 0:
        return [lines]
    chunks = []
    for start in range(0, len(lines), max_lines):
        chunks.append((start + 1, min(len(lines), start + max_lines), lines[start:start + max_lines]))
    return chunks or [(1, 1, [])]


def choose_lexer(file_path):
    from pygments.lexers import TextLexer, get_lexer_for_filename

    try:
        return get_lexer_for_filename(str(file_path))
    except Exception:
        return TextLexer()


def render_code_image(file_path, code_text, font_size, style_name):
    from PIL import Image
    from pygments import highlight
    from pygments.formatters import ImageFormatter

    formatter = ImageFormatter(
        style=style_name,
        line_numbers=True,
        font_name="Consolas",
        font_size=font_size,
        image_pad=18,
        line_pad=3,
        line_number_bg="#1e1e1e",
        line_number_fg="#9aa5b1",
    )
    lexer = choose_lexer(file_path)
    png_bytes = highlight(code_text, lexer, formatter)
    image = Image.open(io.BytesIO(png_bytes)).convert("RGB")
    return image


def fit_image_on_slide(image_width, image_height, max_width_inches, max_height_inches):
    width_inches = image_width / 96.0
    height_inches = image_height / 96.0
    scale = min(max_width_inches / max(width_inches, 0.01), max_height_inches / max(height_inches, 0.01))
    return width_inches * scale, height_inches * scale


def add_title_slide(presentation, title_text, subtitle_text):
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN
    from pptx.util import Inches, Pt

    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    background = slide.background.fill
    background.solid()
    background.fore_color.rgb = RGBColor(20, 24, 32)

    title_box = slide.shapes.add_textbox(Inches(0.8), Inches(1.1), Inches(11.7), Inches(1.2))
    title_frame = title_box.text_frame
    title_frame.clear()
    title_paragraph = title_frame.paragraphs[0]
    title_paragraph.alignment = PP_ALIGN.LEFT
    title_run = title_paragraph.add_run()
    title_run.text = title_text
    title_run.font.size = Pt(28)
    title_run.font.bold = True
    title_run.font.color.rgb = RGBColor(245, 247, 250)

    subtitle_box = slide.shapes.add_textbox(Inches(0.82), Inches(2.15), Inches(11.2), Inches(1.0))
    subtitle_frame = subtitle_box.text_frame
    subtitle_frame.clear()
    subtitle_paragraph = subtitle_frame.paragraphs[0]
    subtitle_paragraph.alignment = PP_ALIGN.LEFT
    subtitle_run = subtitle_paragraph.add_run()
    subtitle_run.text = subtitle_text
    subtitle_run.font.size = Pt(14)
    subtitle_run.font.color.rgb = RGBColor(180, 190, 205)

    footer_box = slide.shapes.add_textbox(Inches(0.82), Inches(6.75), Inches(11.2), Inches(0.5))
    footer_frame = footer_box.text_frame
    footer_paragraph = footer_frame.paragraphs[0]
    footer_run = footer_paragraph.add_run()
    footer_run.text = "Generated {}".format(datetime.now().strftime("%Y-%m-%d %H:%M"))
    footer_run.font.size = Pt(10)
    footer_run.font.color.rgb = RGBColor(135, 145, 160)


def add_code_slide(presentation, relative_path, start_line, end_line, image_path):
    from PIL import Image
    from pptx.dml.color import RGBColor
    from pptx.util import Inches, Pt

    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    background = slide.background.fill
    background.solid()
    background.fore_color.rgb = RGBColor(245, 247, 250)

    header_box = slide.shapes.add_textbox(Inches(0.45), Inches(0.22), Inches(12.1), Inches(0.45))
    header_frame = header_box.text_frame
    header_frame.clear()
    paragraph = header_frame.paragraphs[0]
    run = paragraph.add_run()
    run.text = str(relative_path)
    run.font.size = Pt(22)
    run.font.bold = True
    run.font.color.rgb = RGBColor(26, 32, 44)

    line_box = slide.shapes.add_textbox(Inches(0.47), Inches(0.66), Inches(12.0), Inches(0.28))
    line_frame = line_box.text_frame
    line_frame.clear()
    line_paragraph = line_frame.paragraphs[0]
    line_run = line_paragraph.add_run()
    line_run.text = "Lines {}-{}".format(start_line, end_line)
    line_run.font.size = Pt(11)
    line_run.font.color.rgb = RGBColor(90, 100, 115)

    with Image.open(image_path) as image:
        width_inches, height_inches = fit_image_on_slide(image.width, image.height, 11.8, 6.0)
    left = Inches((13.333 - width_inches) / 2.0)
    top = Inches(1.0)
    slide.shapes.add_picture(str(image_path), left, top, width=Inches(width_inches), height=Inches(height_inches))


def build_presentation(source_paths, output_path, title_text, subtitle_text, project_root, max_lines, font_size, style_name):
    from pptx import Presentation

    presentation = Presentation()
    presentation.slide_width = int(13.333 * 914400)
    presentation.slide_height = int(7.5 * 914400)
    add_title_slide(presentation, title_text, subtitle_text)

    project_root = Path(project_root).expanduser().resolve()
    with tempfile.TemporaryDirectory(prefix="code_ppt_") as temp_dir:
        temp_root = Path(temp_dir)
        slide_index = 0
        for file_path in source_paths:
            relative_path = file_path
            try:
                relative_path = file_path.relative_to(project_root)
            except Exception:
                relative_path = file_path
            file_lines = file_path.read_text(encoding="utf-8", errors="replace").splitlines()
            for start_line, end_line, chunk in chunk_lines(file_lines, max_lines):
                code_text = "\n".join(chunk) if chunk else ""
                image = render_code_image(file_path, code_text, font_size=font_size, style_name=style_name)
                image_path = temp_root / "slide_{:04d}.png".format(slide_index)
                image.save(image_path, format="PNG", optimize=True)
                add_code_slide(presentation, relative_path, start_line, end_line, image_path)
                slide_index += 1

        output_path.parent.mkdir(parents=True, exist_ok=True)
        presentation.save(str(output_path))


def main():
    args = parse_args()
    ensure_dependencies()
    source_paths = expand_sources(args.sources)
    output_path = Path(args.output).expanduser().resolve()
    build_presentation(
        source_paths=source_paths,
        output_path=output_path,
        title_text=args.title,
        subtitle_text=args.subtitle,
        project_root=args.project_root,
        max_lines=args.max_lines,
        font_size=args.font_size,
        style_name=args.style,
    )
    print("Saved PowerPoint to: {}".format(output_path))
    print("Source files included: {}".format(len(source_paths)))


if __name__ == "__main__":
    main()
